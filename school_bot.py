import os
import pickle
import asyncio
import aiofiles
import re
import math
import logging
import json
import ssl
from difflib import get_close_matches
from datetime import datetime, timedelta
import aiohttp
import openpyxl
import pandas as pd
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from vkbottle.bot import Bot, Message
from vkbottle import Keyboard, Text
from dotenv import load_dotenv

os.environ['PYTHONASYNCIODEBUG'] = '0'

# Подхватываем .env, если он лежит рядом со скриптом (достаточно создать
# файл .env с VK_TOKEN=... и MAX_TOKEN=... — больше ничего делать не нужно,
# никаких "set"/"export" в командной строке). Если .env нет — не мешает,
# работаем с тем, что уже задано в системе (например, через Docker/
# docker-compose или set/export).
load_dotenv()

# ─── Токены ───────────────────────────────────────────────────────────────
# Токены ТОЛЬКО из переменных окружения / .env — в коде их нет вообще
# (репозиторий публичный). Проще всего — создать файл .env рядом со
# скриптом с двумя строками:
#   VK_TOKEN=...
#   MAX_TOKEN=...
# Либо задать их как настоящие переменные окружения:
#   Windows (cmd):         set VK_TOKEN=... & set MAX_TOKEN=...
#   Windows (PowerShell):  $env:VK_TOKEN="..."; $env:MAX_TOKEN="..."
#   Linux/macOS:           export VK_TOKEN=... MAX_TOKEN=...
VK_TOKEN = os.environ.get("VK_TOKEN", "")
MAX_TOKEN = os.environ.get("MAX_TOKEN", "")

if not VK_TOKEN or not MAX_TOKEN:
    _missing = [n for n, v in (("VK_TOKEN", VK_TOKEN), ("MAX_TOKEN", MAX_TOKEN)) if not v]
    raise SystemExit(
        "Не заданы переменные окружения: " + ", ".join(_missing) + ". "
        "Укажите их перед запуском (см. комментарий выше в коде)."
    )

# Публичное имя бота в MAX и ссылка на сообщество VK — используются в тексте
# справки (START_TEXT) ниже.
MAX_BOT_USERNAME = "ideya_fiks_7_bot"
VK_COMMUNITY_URL = "https://vk.ru/club242725074"

# Базовый адрес REST API MAX (https://dev.max.ru/docs-api).
# Для чат-ботов и мини-приложений документация рекомендует platform-api2.max.ru
MAX_API_BASE = "https://platform-api2.max.ru"

bot = Bot(token=VK_TOKEN)

logging.basicConfig(level=logging.WARNING)
logging.getLogger("vkbottle").setLevel(logging.WARNING)

# ─── Пути к файлам ────────────────────────────────────────────────────────────

# Актуальный файл расписания (новый формат 2026-2027, см. ниже раздел «Парсинг
# расписания»). Положите xlsx рядом со скриптом под этим именем, либо поменяйте
# путь здесь.
SCHEDULE_FILE       = "schedule.xlsx"
CONSULTATIONS_FILE  = "Transformatsia_raspisania.xlsx"
MENU_1_4_FILE       = "меню 1-4.xlsx"
MENU_5_11_FILE      = "меню 5-11.xlsx"
CHANGES_DIR         = "changes_files"
CHANGES_XLSX_FILE   = os.path.join(CHANGES_DIR, "Замены для чата.xlsx")
CHANGES_DOC_FILE    = os.path.join(CHANGES_DIR, "changes_doc.docx")
CACHE_FILE          = "schedule_cache.pkl"
MENU_1_4_CACHE_FILE = "menu_1_4_cache.pkl"
MENU_5_11_CACHE_FILE= "menu_5_11_cache.pkl"
CHANGES_CACHE_FILE  = os.path.join(CHANGES_DIR, "changes_cache.pkl")
ERROR_LOG_FILE      = "log_error.txt"
STATISTIC_LOG_FILE  = "log_statistic.txt"

os.makedirs(CHANGES_DIR, exist_ok=True)

if not os.path.exists(SCHEDULE_FILE):
    raise FileNotFoundError(f"Файл {SCHEDULE_FILE} не найден!")

# ─── Константы ────────────────────────────────────────────────────────────────

DAY_MAPPING = {
    "пн.": "понедельник", "вт.": "вторник", "ср.": "среда",
    "чтв.": "четверг", "пт.": "пятница", "сб.": "суббота", "вс.": "воскресенье"
}
DAYS = ["понедельник", "вторник", "среда", "четверг", "пятница", "суббота", "воскресенье"]
_DAYS = ["понедельник", "вторник", "среда", "четверг", "пятница"]

MEAL_SCHEDULE = {
    "корпус 1": [
        {"трапеза": "завтрак", "время": "09:00–09:10", "классы": "1 классы, 4 классы"},
        {"трапеза": "завтрак", "время": "09:55–10:05", "классы": "2 классы, льготные, своб. посещение"},
        {"трапеза": "завтрак", "время": "10:50–11:00", "классы": "3 классы, своб. посещение"},
        {"трапеза": "обед",    "время": "11:45–12:05", "классы": "льготные, своб. посещение"},
        {"трапеза": "обед",    "время": "12:50–13:05", "классы": "льготные, своб. посещение"},
        {"трапеза": "полдник", "время": "16:00–16:30", "классы": "своб. посещение"},
    ],
    "корпус 2": [
        {"трапеза": "завтрак", "время": "09:00–09:10", "классы": "1 классы, 2 «Г», 2 «Ж»"},
        {"трапеза": "завтрак", "время": "09:55–10:05", "классы": "2 «Б», 2 «Д», 3 классы"},
        {"трапеза": "завтрак", "время": "10:50–11:00", "классы": "4 классы"},
        {"трапеза": "обед",    "время": "11:45–12:05", "классы": "льготные, своб. посещение"},
        {"трапеза": "обед",    "время": "12:50–13:05", "классы": "своб. посещение"},
        {"трапеза": "полдник", "время": "16:00–16:30", "классы": "своб. посещение"},
    ],
}

START_TEXT = (
    "Добрый день! 👋 Я бот-помощник по расписанию школы.\n"
    "Пишите мне обычным текстом, без специальных команд — понимаю почти любые\n"
    "формулировки и опечатки.\n\n"

    "📅 РАСПИСАНИЕ КЛАССА ИЛИ УЧИТЕЛЯ\n"
    "• «9А понедельник» — расписание класса на день\n"
    "• «5 А вторник» — пробел перед буквой класса не важен\n"
    "• «Иванов вторник» — расписание учителя по фамилии (можно с опечаткой —\n"
    "  найду ближайшее совпадение)\n"
    "• «учителя» — список всех учителей школы и дней, когда у них есть уроки\n\n"

    "🍽 МЕНЮ ПИТАНИЯ\n"
    "• «меню» — пришлю клавиатуру: выбрать классы (1–4 или 5–11) и день\n"
    "• «меню среда» — после выбора классов меню покажется сразу на этот день\n\n"

    "🔁 ИЗМЕНЕНИЯ В РАСПИСАНИИ (ЗАМЕНЫ)\n"
    "• «изменения пятница» — список замен/отмен уроков на пятницу, если они\n"
    "  уже загружены на этот день\n\n"

    "👩‍🏫 ДЛЯ ЗАВУЧЕЙ И ОТВЕТСТВЕННЫХ ЗА РАСПИСАНИЕ\n"
    "Подбор замены прямо в диалоге с ботом — напишите «замены»:\n"
    "1. Выберите день по кнопкам.\n"
    "2. Введите фамилию отсутствующего педагога (можно с опечаткой).\n"
    "3. Нажмите на нужный урок из списка его уроков в этот день.\n"
    "4. Выберите свободного в этот момент учителя из списка (кнопки,\n"
    "   постранично) либо нажмите «Отмена урока».\n"
    "5. Добавьте короткий комментарий к замене или отправьте «-», чтобы\n"
    "   пропустить.\n"
    "6. Повторите для всех нужных уроков дня, затем напишите «конец» —\n"
    "   покажу итог и попрошу подтвердить.\n"
    "7. После подтверждения сам соберу и пришлю готовый xlsx-файл с\n"
    "   заменами — можно распечатать или переслать в учительскую.\n\n"
    "Либо просто пришлите мне готовый файл замен:\n"
    "• xlsx с именем, начинающимся на «Замены», или\n"
    "• docx с «Изменения в расписании на…» в названии —\n"
    "и я сам обновлю общий кэш замен (он общий для VK и MAX: файл, присланный\n"
    "в одном мессенджере, сразу отражается в ответах в другом).\n\n"

    "ℹ️ СЛУЖЕБНОЕ\n"
    "• «начало» / «помощь» / /start — эта справка\n"
    "• «debug_cache» — дамп текущего кэша замен как есть (для проверки\n"
    "  загруженного файла)\n\n"

    "💬 Сообщество бота в VK: https://vk.ru/club242725074"
)

# ─── Глобальные данные (общие для VK и MAX) ───────────────────────────────────

schedule_by_class: dict      = {}
teacher_schedule_list: list  = []
teacher_days: dict           = {}
consultations_by_class: dict = {}
consultations_by_teacher: dict = {}
changes_cache: dict          = {}
menu_1_4_cache: dict         = {}
menu_5_11_cache: dict        = {}

class_requests: dict  = {f"{g}{l}": 0 for g in range(5, 12) for l in "АБВГД"}
teacher_requests: dict = {}
total_requests: int   = 0
menu_requests: int    = 0
current_menu_variant: int = 1

# Состояния диалога замен. Ключ — "<platform>:<uid>", чтобы id пользователей
# VK и MAX (оба — числа) не пересекались друг с другом.
user_states: dict = {}
STATE_IDLE                    = "idle"
STATE_CHOOSE_TEACHER          = "choose_teacher"
STATE_CHOOSE_LESSON           = "choose_lesson"
STATE_CHOOSE_REPLACEMENT      = "choose_replacement"
STATE_CONFIRM_REPLACEMENT     = "confirm_replacement"
STATE_ADD_COMMENT             = "add_comment"
STATE_CONFIRM_CHANGES         = "confirm_changes"


def _state_key(platform: str, uid) -> str:
    return f"{platform}:{uid}"


def get_user_state(platform: str, uid) -> dict:
    return user_states.get(_state_key(platform, uid), {"state": STATE_IDLE, "data": {}})


def set_user_state(platform: str, uid, state: str, data: dict = None) -> None:
    user_states[_state_key(platform, uid)] = {"state": state, "data": data or {}}


def clear_user_state(platform: str, uid) -> None:
    user_states.pop(_state_key(platform, uid), None)


# ─── Утилиты ──────────────────────────────────────────────────────────────────

def get_timestamp() -> str:
    return datetime.now().strftime("%d.%m.%Y %H:%M")


async def log_error(msg: str) -> None:
    ts = get_timestamp()
    async with aiofiles.open(ERROR_LOG_FILE, "a", encoding="utf-8") as f:
        await f.write(f"{ts} - {msg}\n")


async def log_statistic() -> None:
    global class_requests, teacher_requests, total_requests, menu_requests
    date = (datetime.now() - timedelta(days=1)).strftime("%d.%m.%Y")
    async with aiofiles.open(STATISTIC_LOG_FILE, "a", encoding="utf-8") as f:
        await f.write(f"{date}\n")
        for cls in sorted(class_requests):
            await f.write(f"{cls} - {class_requests[cls]}\n")
        await f.write("Педагоги:\n")
        for t in sorted(teacher_requests):
            await f.write(f"{t} - {teacher_requests[t]}\n")
        await f.write(f"Всего - {total_requests}\nЗапросы меню - {menu_requests}\n\n")
    class_requests    = {k: 0 for k in class_requests}
    teacher_requests.clear()
    total_requests = 0
    menu_requests  = 0


def get_current_menu_variant() -> int:
    global current_menu_variant
    base = datetime(2025, 5, 13)
    diff = (datetime.now() - base).days
    current_menu_variant = 2 if (diff // 7) % 2 == 0 else 1
    return current_menu_variant


def _sort_key(item):
    k = item[0]
    try:
        return (int(k.split("-")[0]), k)
    except (ValueError, AttributeError):
        return (9999, k)


def _month_name(month: int) -> str:
    """Возвращает название месяца в родительном падеже."""
    names = {
        1: "января", 2: "февраля", 3: "марта", 4: "апреля",
        5: "мая", 6: "июня", 7: "июля", 8: "августа",
        9: "сентября", 10: "октября", 11: "ноября", 12: "декабря"
    }
    return names.get(month, str(month))


def _extract_date(text: str) -> str | None:
    """
    Пытается извлечь дату из строки.
    Поддерживает форматы:
      - «14 апреля», «14 апреля 2025»
      - «14.04», «14.04.2025», «14.04.25»
    Возвращает строку вида «14 апреля» или None.
    """
    m = re.search(r"\b(\d{1,2})\.(\d{1,2})(?:\.\d{2,4})?\b", text)
    if m:
        try:
            day_n = int(m.group(1))
            mon_n = int(m.group(2))
            if 1 <= day_n <= 31 and 1 <= mon_n <= 12:
                return f"{day_n} {_month_name(mon_n)}"
        except ValueError:
            pass

    month_names_re = (
        r"январ[яь]|феврал[яь]|март[а]?|апрел[яь]|ма[йя]|июн[яь]|"
        r"июл[яь]|август[а]?|сентябр[яь]|октябр[яь]|ноябр[яь]|декабр[яь]"
    )
    m2 = re.search(r"\b(\d{1,2})\s+(" + month_names_re + r")\b", text, re.IGNORECASE)
    if m2:
        return f"{int(m2.group(1))} {m2.group(2).lower()}"

    return None


def _cab_display(v) -> str:
    """Приводит значение ячейки «каб.» к отображаемой строке."""
    if v is None:
        return ""
    if isinstance(v, float) and math.isnan(v):
        return ""
    if isinstance(v, str):
        return v.replace("\n", " ").strip()
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return repr(v)
    return str(v).strip()


def _split_cab(v, n: int) -> list:
    """
    Excel часто «склеивает» два кабинета при подгрупповых занятиях
    (например, «102.104» превращается в float 102.104). Разбираем такие
    значения обратно на список кабинетов длиной n. Если значение не похоже
    на «склейку» — просто повторяем одну и ту же подпись n раз.
    """
    if isinstance(v, float) and not math.isnan(v):
        parts = repr(v).split(".")
        if len(parts) == 2 and parts[1]:
            rooms = list(parts)
            while len(rooms) < n:
                rooms.append(rooms[-1])
            return rooms[:n]
    label = _cab_display(v)
    return [label] * n


# ─── Фоновые задачи ───────────────────────────────────────────────────────────

async def switch_menu_variant():
    global current_menu_variant
    while True:
        now = datetime.now()
        days_to_sat = (5 - now.weekday() + 7) % 7 or 7
        next_sat = (now + timedelta(days=days_to_sat)).replace(hour=23, minute=0, second=0, microsecond=0)
        await asyncio.sleep((next_sat - now).total_seconds())
        current_menu_variant = 2 if current_menu_variant == 1 else 1
        print(f"{get_timestamp()} - Вариант меню: {current_menu_variant}")
        await asyncio.sleep(60)


async def check_midnight():
    while True:
        now = datetime.now()
        next_mn = (now + timedelta(days=1)).replace(hour=0, minute=0, second=0, microsecond=0)
        await asyncio.sleep((next_mn - now).total_seconds())
        await log_statistic()
        await asyncio.sleep(1)


# ─── Парсинг расписания (новый формат 2026-2027) ──────────────────────────────
#
# Новый файл устроен иначе, чем старый:
#   - один лист "Расписание 1й корпус" (без деления на «1/2 неделю»);
#   - блоки по 12 строк на класс: строка с названием класса (например «5 А»),
#     строка заголовков («урок»/день/«учитель»/«каб.» x5), 8 строк уроков;
#   - для каждого дня — 4 колонки: урок, предмет, учитель, кабинет,
#     смещение колонки предмета = 1 + day_idx*4 (0-индексация pandas), как и
#     раньше;
#   - при делении класса на подгруппы (труд, английский, информатика и т.п.)
#     учитель записан как две строки через перенос "\n" в одной ячейке,
#     а Excel автоматически "склеивает" два номера кабинета в один float
#     (например "102.104" -> 102.104). Разбираем это в _split_cab().
#
# Функция сама находит все листы, чьё имя начинается с «Расписание» — если в
# будущем добавится лист по 2 корпусу или ещё одна неделя, менять код не
# придётся.

def load_or_parse_schedule():
    global schedule_by_class, teacher_schedule_list, teacher_days

    # Кэш валиден только если он новее самого файла расписания — иначе после
    # замены xlsx бот продолжал бы отвечать по старому кэшу.
    if os.path.exists(CACHE_FILE) and os.path.getmtime(CACHE_FILE) >= os.path.getmtime(SCHEDULE_FILE):
        try:
            with open(CACHE_FILE, "rb") as f:
                data = pickle.load(f)
            schedule_by_class     = data["schedule_by_class"]
            teacher_schedule_list = data["teacher_schedule_list"]
            teacher_days          = data["teacher_days"]
            print(f"{get_timestamp()} - Кэш расписания загружен ({len(schedule_by_class)} классов)")
            return
        except Exception as e:
            logging.error(f"Ошибка загрузки кэша: {e}")

    schedule_by_class = {}
    teacher_schedule_list = []
    teacher_days = {}

    try:
        excel = pd.ExcelFile(SCHEDULE_FILE)
    except Exception as e:
        logging.error(f"Не удалось открыть {SCHEDULE_FILE}: {e}")
        return

    sheet_names = [s for s in excel.sheet_names if s.strip().lower().startswith("расписание")]
    if not sheet_names:
        logging.error(
            f"В файле {SCHEDULE_FILE} не найдено ни одного листа с расписанием "
            f"(имя листа должно начинаться со слова «Расписание»)"
        )

    for sheet_name in sheet_names:
        try:
            df = pd.read_excel(excel, sheet_name=sheet_name, header=None)
        except Exception as e:
            logging.error(f"Ошибка чтения '{sheet_name}': {e}")
            continue

        row = 0
        while row < len(df):
            cell = str(df.iloc[row, 1]).strip()
            if re.match(r"^\d{1,2}\s*[А-Я]$", cell, re.IGNORECASE):
                class_name = "".join(cell.split()).upper()
                current_row = row + 2  # пропускаем строку заголовков
                while current_row < len(df) and str(df.iloc[current_row, 0]).strip().isdigit():
                    lesson_num = str(df.iloc[current_row, 0]).strip()
                    for day_idx, day in enumerate(_DAYS):
                        cl = 1 + day_idx * 4
                        subj_raw    = df.iloc[current_row, cl]
                        teacher_raw = df.iloc[current_row, cl + 1]
                        cab_raw     = df.iloc[current_row, cl + 2]

                        if isinstance(subj_raw, float) and math.isnan(subj_raw):
                            continue
                        subj = str(subj_raw).strip()
                        if not subj or subj.lower() == "nan":
                            continue

                        teacher_str = "" if (isinstance(teacher_raw, float) and math.isnan(teacher_raw)) \
                            else str(teacher_raw).strip()
                        if not teacher_str or teacher_str.lower() == "nan":
                            continue

                        if "\n" in teacher_str:
                            # Подгруппа: два (реже больше) учителя на один и тот же предмет
                            teachers = [t.strip() for t in teacher_str.split("\n") if t.strip()]
                            cabs = _split_cab(cab_raw, len(teachers))
                            for tf, cab in zip(teachers, cabs):
                                if not tf:
                                    continue
                                surn = tf.split()[0].capitalize() if " " in tf else tf.capitalize()
                                schedule_by_class.setdefault(class_name, {}).setdefault(day, {}) \
                                    .setdefault(lesson_num, []).append(f"{subj} ({tf}), {cab} каб.")
                                teacher_schedule_list.append({
                                    "surname": surn, "day": day, "lesson_num": lesson_num,
                                    "info": f"{class_name} - {subj}, {cab} каб."
                                })
                                teacher_days.setdefault(surn, set()).add(day)
                        else:
                            tf  = teacher_str
                            cab = _cab_display(cab_raw)
                            surn = tf.split()[0].capitalize() if " " in tf else tf.capitalize()
                            schedule_by_class.setdefault(class_name, {}).setdefault(day, {}) \
                                .setdefault(lesson_num, []).append(f"{subj} ({tf}), {cab} каб.")
                            teacher_schedule_list.append({
                                "surname": surn, "day": day, "lesson_num": lesson_num,
                                "info": f"{class_name} - {subj}, {cab} каб."
                            })
                            teacher_days.setdefault(surn, set()).add(day)
                    current_row += 1
                row = current_row
            else:
                row += 1

    try:
        with open(CACHE_FILE, "wb") as f:
            pickle.dump({"schedule_by_class": schedule_by_class,
                         "teacher_schedule_list": teacher_schedule_list,
                         "teacher_days": teacher_days}, f)
        print(f"{get_timestamp()} - Кэш расписания сохранён ({len(schedule_by_class)} классов, "
              f"{len(teacher_days)} учителей)")
    except Exception as e:
        logging.error(f"Ошибка сохранения кэша: {e}")


# ─── Парсинг меню (без изменений) ──────────────────────────────────────────────

def load_or_parse_menu():
    global menu_1_4_cache, menu_5_11_cache

    def _load(cache_f, src_f):
        if os.path.exists(cache_f) and os.path.exists(src_f):
            if os.path.getmtime(src_f) <= os.path.getmtime(cache_f):
                with open(cache_f, "rb") as f:
                    return pickle.load(f)
        if not os.path.exists(src_f):
            return {}
        data = parse_menu_file(src_f)
        with open(cache_f, "wb") as f:
            pickle.dump(data, f)
        return data

    menu_1_4_cache  = _load(MENU_1_4_CACHE_FILE,  MENU_1_4_FILE)
    menu_5_11_cache = _load(MENU_5_11_CACHE_FILE, MENU_5_11_FILE)


def parse_menu_file(file_path: str) -> dict:
    try:
        wb    = openpyxl.load_workbook(file_path)
        sheet = wb.active
        is14  = "1-4" in file_path.lower()
        menu_data: dict = {}

        day_configs = [
            {"start_row": 5,  "day_name": "понедельник", "meals": [{"type":"завтрак","row":6,"count":5},{"type":"обед","row":13,"count":6},{"type":"полдник","row":21,"count":5}]},
            {"start_row": 29, "day_name": "вторник",     "meals": [{"type":"завтрак","row":30,"count":5},{"type":"обед","row":37,"count":7},{"type":"полдник","row":46,"count":6}]},
            {"start_row": 55, "day_name": "среда",       "meals": [{"type":"завтрак","row":56,"count":5},{"type":"обед","row":63,"count":7},{"type":"полдник","row":72,"count":4}]},
            {"start_row": 79, "day_name": "четверг",     "meals": [{"type":"завтрак","row":80,"count":4},{"type":"обед","row":86,"count":6},{"type":"полдник","row":94,"count":4}]},
            {"start_row": 101,"day_name": "пятница",     "meals": [{"type":"завтрак","row":102,"count":5},{"type":"обед","row":109,"count":6},{"type":"полдник","row":117,"count":5}]},
            {"start_row": 126,"day_name": "понедельник", "meals": [{"type":"завтрак","row":127,"count":5},{"type":"обед","row":134,"count":7},{"type":"полдник","row":143,"count":5}]},
            {"start_row": 151,"day_name": "вторник",     "meals": [{"type":"завтрак","row":152,"count":5},{"type":"обед","row":159,"count":5},{"type":"полдник","row":168,"count":4}]},
            {"start_row": 175,"day_name": "среда",       "meals": [{"type":"завтрак","row":176,"count":5},{"type":"обед","row":183,"count":7},{"type":"полдник","row":192,"count":6}]},
            {"start_row": 201,"day_name": "четверг",     "meals": [{"type":"завтрак","row":202,"count":4 if is14 else 5},{"type":"обед","row":208 if is14 else 209,"count":6},{"type":"полдник","row":216 if is14 else 217,"count":4}]},
            {"start_row": 223 if is14 else 224,"day_name":"пятница","meals":[{"type":"завтрак","row":224 if is14 else 225,"count":5},{"type":"обед","row":231 if is14 else 232,"count":6},{"type":"полдник","row":239 if is14 else 240,"count":5}]},
        ]

        for cfg in day_configs:
            day = cfg["day_name"]
            dc  = sheet.cell(row=cfg["start_row"], column=1).value
            if not dc or not isinstance(dc, str) or not re.match(r".*?\(.+?\)", dc):
                continue
            meal_data = {}
            for m in cfg["meals"]:
                mc = sheet.cell(row=m["row"], column=1).value
                if not mc or m["type"] not in str(mc).lower():
                    continue
                dishes = []
                for i in range(m["count"]):
                    ri   = m["row"] + 1 + i
                    num  = sheet.cell(row=ri, column=1).value
                    dish = sheet.cell(row=ri, column=2).value
                    if num and dish and str(num).strip().isdigit():
                        dishes.append((int(num), str(dish).strip()))
                if dishes:
                    meal_data[m["type"]] = dishes
            if meal_data:
                menu_data.setdefault(day, []).append(meal_data)
        return menu_data
    except Exception as e:
        logging.error(f"Ошибка парсинга меню '{file_path}': {e}")
        return {}


# ─── Парсинг изменений (без изменений) ────────────────────────────────────────

async def parse_doc_changes(file_path: str):
    try:
        doc = Document(file_path)
    except PackageNotFoundError:
        await log_error(f"'{file_path}' не является Word-документом")
        return None

    day = None
    change_date = None

    for para in doc.paragraphs:
        txt = para.text.strip()
        txt_low = txt.lower()
        if "изменения в расписании" in txt_low:
            for d in DAYS:
                if d in txt_low:
                    day = d
                    break
            change_date = _extract_date(txt_low)
            if day:
                break

    if not day:
        await log_error(f"День не найден в '{file_path}'")
        return None

    table_data = None
    for table in doc.tables:
        td = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if len(td) > 1 and len(td[0]) > 1:
            table_data = td
            break
    if not table_data:
        await log_error(f"Таблица не найдена в '{file_path}'")
        return None

    changes: dict = {}
    classes = table_data[0][1:]
    for row in table_data[1:]:
        if not row[0] or not row[0].isdigit():
            continue
        ln = row[0]
        for j, change in enumerate(row[1:], 1):
            cn = "".join(classes[j - 1].split()).upper()
            if change and change != "-":
                changes.setdefault(cn, {})[ln] = change

    if not changes:
        await log_error(f"Изменения не найдены в '{file_path}'")
        return None

    return {day: {"changes": changes, "date": change_date}}


async def parse_xlsx_changes(file_path: str):
    try:
        wb    = openpyxl.load_workbook(file_path)
        sheet = wb.worksheets[-1]
        day_cell = sheet["B2"].value
        if not day_cell:
            await log_error(f"B2 пуста в '{file_path}'")
            return None

        day = None
        change_date = None

        if isinstance(day_cell, datetime):
            change_date = f"{day_cell.day} {_month_name(day_cell.month)}"
            day_str = day_cell.strftime("%d.%m.%Y")
        else:
            day_str = str(day_cell).strip().lower()
            change_date = _extract_date(day_str)

        for d in DAYS:
            if d in day_str:
                day = d
                break

        if not day:
            for fmt in ("%d.%m.%Y", "%d.%m.%y", "%d.%m"):
                try:
                    parsed_dt = datetime.strptime(day_str.strip(), fmt)
                    day = DAYS[parsed_dt.weekday()] if parsed_dt.weekday() < 7 else None
                    break
                except ValueError:
                    continue

        if not day or day not in DAYS:
            await log_error(f"Не удалось извлечь день из B2='{day_cell}'")
            return None

        changes: dict = {}
        for row in sheet.iter_rows(min_row=6, values_only=True):
            if not row or len(row) < 7:
                continue
            ln_orig   = str(row[1]).strip() if row[1] else ""
            cs        = str(row[2]).strip() if row[2] else ""
            ln_repl   = str(row[4]).strip() if row[4] else ""
            sr        = str(row[5]).strip() if row[5] else ""
            note      = str(row[6]).strip() if row[6] else ""

            if not cs:
                continue
            parts = cs.split()
            if len(parts) < 2:
                continue

            cn = None; end_idx = 1
            if re.match(r"^\d{1,2}[А-Я]$", parts[0], re.IGNORECASE):
                cn = "".join(parts[0].split()).upper()
            if not cn and len(parts) > 2:
                cand = " ".join(parts[:2])
                if re.match(r"^\d{1,2}\s*[А-Я]$", cand):
                    cn = "".join(cand.split()).upper(); end_idx = 2
            if not cn or not re.match(r"^\d{1,2}[А-Я]$", cn):
                continue

            orig = " ".join(parts[end_idx:])
            if not orig:
                continue
            ln = ln_repl or ln_orig
            if not ln:
                continue
            try:    ln_key = str(int(ln))
            except: ln_key = ln

            replacement = ""
            if sr:
                rp = sr.split(); ri = 1
                if rp and re.match(r"^\d{1,2}[А-Я]$", rp[0], re.IGNORECASE):
                    rc = "".join(rp[0].split()).upper()
                    replacement = f"{rc} {' '.join(rp[ri:])}"
                else:
                    replacement = sr

            if "отмен" in note.lower():
                ct = f"{orig} отменён ({note})"
            elif replacement:
                ct = f"{orig} заменён на {replacement} ({note})"
            else:
                ct = f"{orig} ({note})"

            changes.setdefault(cn, {})[ln_key] = ct

            if note:
                fw = note.split()[0].capitalize()
                if fw in teacher_days:
                    teacher_ct = (
                        f"{cn} - {orig} отменён ({note})" if "отмен" in note.lower()
                        else f"{cn} - {orig} заменён на {replacement} ({note})" if replacement
                        else f"{cn} - {orig} ({note})"
                    )
                    changes.setdefault(fw, {})[ln_key] = teacher_ct

        if not changes:
            await log_error(f"Изменения не найдены в '{file_path}'")
            return None

        return {day: {"changes": changes, "date": change_date}}

    except Exception as e:
        await log_error(f"Ошибка XLSX '{file_path}': {e}")
        return None


async def update_changes_cache():
    global changes_cache
    changes_cache = {}

    if os.path.exists(CHANGES_XLSX_FILE):
        ch = await parse_xlsx_changes(CHANGES_XLSX_FILE)
        if ch:
            for d, v in ch.items():
                changes_cache.setdefault(d, {"changes": {}, "date": None})
                changes_cache[d]["changes"].update(v["changes"])
                if v.get("date"):
                    changes_cache[d]["date"] = v["date"]

    if os.path.exists(CHANGES_DOC_FILE):
        ch = await parse_doc_changes(CHANGES_DOC_FILE)
        if ch:
            for d, v in ch.items():
                changes_cache.setdefault(d, {"changes": {}, "date": None})
                changes_cache[d]["changes"].update(v["changes"])
                if v.get("date"):
                    changes_cache[d]["date"] = v["date"]

    if changes_cache:
        with open(CHANGES_CACHE_FILE, "wb") as f:
            pickle.dump(changes_cache, f)
    elif os.path.exists(CHANGES_CACHE_FILE):
        os.remove(CHANGES_CACHE_FILE)


async def get_changes_for_class(cn: str, day: str) -> str:
    if not changes_cache or day not in changes_cache:
        return f"\n\nИзменений на {day.capitalize()} пока нет."
    entry = changes_cache[day]
    ch = entry["changes"].get(cn, {})
    if not ch:
        return "\n\nИзменений нет."
    date_str = entry.get("date")
    header = f"\n\nИзменения{f' на {date_str}' if date_str else ''}:\n"
    return header + "".join(f"Урок {k}: {v}\n" for k, v in sorted(ch.items(), key=_sort_key))


async def get_changes_for_teacher(surn: str, day: str) -> str:
    if not changes_cache or day not in changes_cache:
        return f"\n\nИзменений на {day.capitalize()} пока нет."
    entry = changes_cache[day]
    ch = entry["changes"].get(surn, {})
    if not ch:
        return "\n\nИзменений нет."
    date_str = entry.get("date")
    header = f"\n\nИзменения{f' на {date_str}' if date_str else ''}:\n"
    return header + "".join(f"Урок {k}: {v}\n" for k, v in sorted(ch.items(), key=_sort_key))


# ─── Клавиатуры: единое (платформонезависимое) представление ────────────────
#
# Каждая клавиатура — список строк, строка — список кнопок:
#   {"text": "...", "payload": {...}}
# Рендер под конкретную платформу — в vk_kb_from_neutral() / max_kb_from_neutral().

def nkb_days_changes():
    return [
        [{"text": "Понедельник", "payload": {"cmd": "changes_day", "day": "понедельник"}},
         {"text": "Вторник",     "payload": {"cmd": "changes_day", "day": "вторник"}}],
        [{"text": "Среда",       "payload": {"cmd": "changes_day", "day": "среда"}},
         {"text": "Четверг",     "payload": {"cmd": "changes_day", "day": "четверг"}}],
        [{"text": "Пятница",     "payload": {"cmd": "changes_day", "day": "пятница"}}],
    ]


def nkb_menu_main():
    return [
        [{"text": "📋 Расписание трапез", "payload": {"cmd": "menu_schedule"}}],
        [{"text": "Понедельник", "payload": {"cmd": "menu_day", "day": "понедельник"}},
         {"text": "Вторник",     "payload": {"cmd": "menu_day", "day": "вторник"}}],
        [{"text": "Среда",       "payload": {"cmd": "menu_day", "day": "среда"}},
         {"text": "Четверг",     "payload": {"cmd": "menu_day", "day": "четверг"}}],
        [{"text": "Пятница",     "payload": {"cmd": "menu_day", "day": "пятница"}}],
    ]


def nkb_menu_classes(day: str):
    return [
        [{"text": "1–4 классы",  "payload": {"cmd": "menu_classes", "day": day, "cl": "1-4"}}],
        [{"text": "5–11 классы", "payload": {"cmd": "menu_classes", "day": day, "cl": "5-11"}}],
    ]


def nkb_lessons(nums: list):
    rows = []
    row  = []
    for n in nums:
        row.append({"text": f"Урок {n}", "payload": {"cmd": "lesson", "num": n}})
        if len(row) == 3:
            rows.append(row); row = []
    if row:
        rows.append(row)
    rows.append([{"text": "✅ Готово", "payload": {"cmd": "lessons_done"}}])
    return rows


TEACHERS_PER_PAGE = 6


def nkb_teachers_page(avail: list, page: int):
    total = len(avail)
    total_pages = max(1, (total + TEACHERS_PER_PAGE - 1) // TEACHERS_PER_PAGE)
    page = max(0, min(page, total_pages - 1))
    start = page * TEACHERS_PER_PAGE
    chunk = avail[start: start + TEACHERS_PER_PAGE]

    has_prev = page > 0
    has_next = page < total_pages - 1

    rows = [[{"text": "❌ Отмена урока", "payload": {"cmd": "cancel_lesson"}}]]

    for i in range(0, len(chunk), 2):
        r = [{"text": chunk[i][:35], "payload": {"cmd": "replace", "teacher": chunk[i]}}]
        if i + 1 < len(chunk):
            r.append({"text": chunk[i + 1][:35], "payload": {"cmd": "replace", "teacher": chunk[i + 1]}})
        rows.append(r)

    nav = []
    if has_prev:
        nav.append({"text": "◀ Назад", "payload": {"cmd": "repl_page", "page": page - 1}})
    if has_next:
        nav.append({"text": "Далее ▶", "payload": {"cmd": "repl_page", "page": page + 1}})
    if nav:
        rows.append(nav)

    return rows


def nkb_yes_no(yes_cmd: str = "confirm_yes", no_cmd: str = "confirm_no"):
    return [[{"text": "✅ Да", "payload": {"cmd": yes_cmd}},
             {"text": "❌ Нет", "payload": {"cmd": no_cmd}}]]


def nkb_confirm():
    return nkb_yes_no("confirm_yes", "confirm_no")


def nkb_final_confirm():
    return nkb_yes_no("confirm_final", "cancel_final")


def vk_kb_from_neutral(neutral):
    if not neutral:
        return None
    kb = Keyboard(inline=True)
    for i, row in enumerate(neutral):
        if i > 0:
            kb.row()
        for btn in row:
            kb.add(Text(btn["text"], payload=btn["payload"]))
    return kb.get_json()


def max_kb_from_neutral(neutral):
    """Возвращает список attachments для POST /messages (MAX)."""
    if not neutral:
        return None
    buttons = []
    for row in neutral:
        btn_row = []
        for btn in row:
            btn_row.append({
                "type": "callback",
                "text": btn["text"][:128],
                "payload": json.dumps(btn["payload"], ensure_ascii=False)[:1000],
            })
        buttons.append(btn_row)
    return [{"type": "inline_keyboard", "payload": {"buttons": buttons}}]


def parse_payload(message: Message) -> dict:
    try:
        raw = message.payload
        if not raw:
            return {}
        if isinstance(raw, str):
            return json.loads(raw)
        if isinstance(raw, dict):
            return raw
    except Exception:
        pass
    return {}


# ─── Отправка документа через VK API ─────────────────────────────────────────

async def vk_send_document(peer_id: int, file_path: str, title: str) -> bool:
    try:
        upload_server = await bot.api.docs.get_messages_upload_server(peer_id=peer_id, type="doc")
        upload_url    = upload_server.upload_url
        async with aiohttp.ClientSession() as session:
            with open(file_path, "rb") as fp:
                form = aiohttp.FormData()
                form.add_field("file", fp, filename=title)
                async with session.post(upload_url, data=form) as resp:
                    result = await resp.json(content_type=None)
        saved = await bot.api.docs.save(file=result["file"], title=title)
        att   = f"doc{saved.doc.owner_id}_{saved.doc.id}"
        await bot.api.messages.send(peer_id=peer_id, attachment=att,
                                    message=f"Файл с заменами: {title}", random_id=0)
        return True
    except Exception as e:
        await log_error(f"Ошибка отправки файла (VK): {e}")
        return False


# ─── MAX: низкоуровневый REST-клиент ─────────────────────────────────────────
# Документация: https://dev.max.ru/docs-api
# Авторизация — заголовком Authorization: <access_token> (без "Bearer").
#
# platform-api2.max.ru использует сертификат, выпущенный Национальным
# удостоверяющим центром Минцифры (Russian Trusted CA), которого нет в
# стандартном системном хранилище доверенных сертификатов. Без него запросы
# падают с SSLCertVerificationError. Ниже — разовая (с кэшем в файл)
# загрузка корневого и промежуточного сертификатов с gu-st.ru и построение
# на их основе SSL-контекста для всех запросов к MAX.

RU_TRUSTED_CA_FILE = "russian_trusted_ca.pem"
RU_TRUSTED_CA_URLS = [
    "https://gu-st.ru/content/lending/russian_trusted_root_ca_pem.crt",
    "https://gu-st.ru/content/lending/russian_trusted_sub_ca_pem.crt",
]

_max_ssl_context: ssl.SSLContext | None = None


async def ensure_ru_trusted_ca() -> None:
    """Скачивает сертификаты Минцифры один раз и кэширует их в файл рядом
    со скриптом. Если скачать не удалось (например, нет доступа к
    gu-st.ru) — просто логируем и работаем с тем, что есть (системным
    хранилищем); при следующем запуске попытка повторится."""
    if os.path.exists(RU_TRUSTED_CA_FILE):
        return
    chunks = []
    try:
        async with aiohttp.ClientSession() as session:
            for url in RU_TRUSTED_CA_URLS:
                try:
                    async with session.get(url, timeout=aiohttp.ClientTimeout(total=20)) as resp:
                        if resp.status == 200:
                            chunks.append((await resp.text()).strip())
                        else:
                            await log_error(f"Не удалось скачать {url}: HTTP {resp.status}")
                except Exception as e:
                    await log_error(f"Не удалось скачать {url}: {e}")
    except Exception as e:
        await log_error(f"Ошибка при загрузке сертификатов Минцифры: {e}")

    if chunks:
        try:
            async with aiofiles.open(RU_TRUSTED_CA_FILE, "w", encoding="utf-8") as f:
                await f.write("\n".join(chunks) + "\n")
            print(f"{get_timestamp()} - Сертификаты Минцифры сохранены в {RU_TRUSTED_CA_FILE}")
        except Exception as e:
            await log_error(f"Не удалось сохранить {RU_TRUSTED_CA_FILE}: {e}")


def get_max_ssl_context() -> ssl.SSLContext:
    """SSL-контекст для запросов к platform-api2.max.ru: системное хранилище
    доверенных сертификатов + сертификаты Минцифры (если файл уже скачан
    функцией ensure_ru_trusted_ca(); можно также положить его вручную —
    файл russian_trusted_ca.pem рядом со скриптом, объединив оба .crt с
    https://www.gosuslugi.ru/crt в один PEM-файл)."""
    global _max_ssl_context
    if _max_ssl_context is not None:
        return _max_ssl_context
    ctx = ssl.create_default_context()
    if os.path.exists(RU_TRUSTED_CA_FILE):
        try:
            ctx.load_verify_locations(cafile=RU_TRUSTED_CA_FILE)
        except Exception as e:
            logging.error(f"Не удалось подключить {RU_TRUSTED_CA_FILE}: {e}")
    _max_ssl_context = ctx
    return ctx


def max_client_session() -> aiohttp.ClientSession:
    """aiohttp.ClientSession с SSL-контекстом, знающим про сертификаты
    Минцифры — использовать для всех запросов к MAX API."""
    connector = aiohttp.TCPConnector(ssl=get_max_ssl_context())
    return aiohttp.ClientSession(connector=connector)


async def max_api_request(session, method, path, params=None, json_body=None):
    url = MAX_API_BASE + path
    headers = {"Authorization": MAX_TOKEN}
    try:
        async with session.request(method, url, params=params, json=json_body,
                                    headers=headers,
                                    timeout=aiohttp.ClientTimeout(total=95)) as resp:
            try:
                data = await resp.json(content_type=None)
            except Exception:
                data = {}
            return resp.status, data
    except Exception as e:
        await log_error(f"MAX API запрос упал ({method} {path}): {e}")
        return 0, {}


async def max_send_message(text: str, user_id=None, chat_id=None, attachments=None) -> dict:
    params = {}
    if user_id is not None:
        params["user_id"] = user_id
    if chat_id is not None:
        params["chat_id"] = chat_id
    body = {"text": (text or "")[:4000]}
    if attachments:
        body["attachments"] = attachments
    async with max_client_session() as session:
        status, data = await max_api_request(session, "POST", "/messages", params=params, json_body=body)
        if status != 200:
            await log_error(f"MAX sendMessage ошибка {status}: {data}")
        return data


async def max_answer_callback(callback_id: str) -> None:
    """Закрывает "часики" на кнопке после обработки нажатия."""
    if not callback_id:
        return
    async with max_client_session() as session:
        status, data = await max_api_request(session, "POST", "/answers",
                                              params={"callback_id": callback_id}, json_body={})
        if status != 200:
            await log_error(f"MAX answerCallback ошибка {status}: {data}")


async def max_upload_and_send_document(peer_id: int, is_chat: bool, file_path: str, title: str) -> bool:
    """
    Загружает файл в MAX (POST /uploads -> multipart на выданный URL) и
    отправляет его вложением в сообщение.
    """
    try:
        async with max_client_session() as session:
            status, up = await max_api_request(session, "POST", "/uploads", params={"type": "file"})
            if status != 200 or not up.get("url"):
                await log_error(f"MAX getUploadUrl ошибка {status}: {up}")
                return False
            upload_url = up["url"]

            with open(file_path, "rb") as fp:
                form = aiohttp.FormData()
                form.add_field("data", fp.read(), filename=os.path.basename(title))
                async with session.post(upload_url, data=form,
                                        headers={"Authorization": MAX_TOKEN}) as resp:
                    try:
                        upload_result = await resp.json(content_type=None)
                    except Exception:
                        upload_result = {}

            token = upload_result.get("token") or upload_result.get("id")
            if not token:
                await log_error(f"MAX: не удалось получить токен файла после загрузки: {upload_result}")
                return False

            attachments = [{"type": "file", "payload": {"token": token}}]
            params = {"chat_id": peer_id} if is_chat else {"user_id": peer_id}
            body = {"text": f"Файл с заменами: {title}", "attachments": attachments}
            status, data = await max_api_request(session, "POST", "/messages", params=params, json_body=body)
            if status != 200:
                await log_error(f"MAX sendDocument ошибка {status}: {data}")
                return False
            return True
    except Exception as e:
        await log_error(f"MAX sendDocument исключение: {e}")
        return False


# ─── Responder: единый интерфейс отправки для VK и MAX ───────────────────────

class Responder:
    platform: str
    uid: int

    async def answer(self, text: str, keyboard=None):
        raise NotImplementedError

    async def send_document(self, path: str, title: str) -> bool:
        raise NotImplementedError


class VKResponder(Responder):
    def __init__(self, message: Message):
        self.platform = "vk"
        self.uid = message.from_id
        self.peer_id = message.peer_id
        self._message = message

    async def answer(self, text: str, keyboard=None):
        await self._message.answer(text, keyboard=vk_kb_from_neutral(keyboard))

    async def send_document(self, path: str, title: str) -> bool:
        return await vk_send_document(self.peer_id, path, title)


class MAXResponder(Responder):
    def __init__(self, uid: int, chat_id=None, callback_id=None):
        self.platform = "max"
        self.uid = uid
        self.chat_id = chat_id
        self.is_chat = chat_id is not None
        self.peer_id = chat_id if chat_id is not None else uid
        self._callback_id = callback_id

    async def answer(self, text: str, keyboard=None):
        attachments = max_kb_from_neutral(keyboard)
        await max_send_message(
            text,
            user_id=None if self.is_chat else self.uid,
            chat_id=self.chat_id if self.is_chat else None,
            attachments=attachments,
        )
        if self._callback_id:
            await max_answer_callback(self._callback_id)
            self._callback_id = None  # отвечаем на конкретное нажатие один раз

    async def send_document(self, path: str, title: str) -> bool:
        return await max_upload_and_send_document(self.peer_id, self.is_chat, path, title)


# ─── Приём файлов замен (общий код для VK и MAX) ─────────────────────────────

async def handle_incoming_file(get_bytes, fname: str, responder: Responder):
    """
    get_bytes — awaitable, возвращающая bytes содержимого файла.
    Логика идентична для VK и MAX: файл сохраняется в общее хранилище,
    после чего пересобирается общий кэш замен changes_cache, которым
    пользуются запросы из обеих платформ.
    """
    flower = fname.lower()
    data = await get_bytes

    if flower.endswith(".xlsx"):
        if not flower.startswith("замены"):
            await responder.answer("Ожидается файл, начинающийся с «Замены» (.xlsx).")
            return
        if os.path.exists(CHANGES_XLSX_FILE):
            os.remove(CHANGES_XLSX_FILE)
        async with aiofiles.open(CHANGES_XLSX_FILE, "wb") as f:
            await f.write(data)
        await responder.answer(f"Файл «{fname}» сохранён.")
        await update_changes_cache()
        await responder.answer("Кэш изменений обновлён ✅")

    elif flower.endswith((".doc", ".docx")):
        if "измененияврасписаниина" not in flower.replace(" ", "").replace("_", ""):
            await responder.answer("Ожидается файл «Изменения в расписании на...» (.docx).")
            return
        if os.path.exists(CHANGES_DOC_FILE):
            os.remove(CHANGES_DOC_FILE)
        async with aiofiles.open(CHANGES_DOC_FILE, "wb") as f:
            await f.write(data)
        await responder.answer(f"Файл «{fname}» сохранён.")
        await update_changes_cache()
        await responder.answer("Кэш изменений обновлён ✅")

    else:
        await responder.answer("Поддерживаются только .xlsx (Замены) и .docx (Изменения).")


# ─── Генерация xlsx с заменами (без изменений) ────────────────────────────────

async def generate_xlsx_changes(day: str, changes: dict, output_path: str,
                                temp_rows: list = None) -> bool:
    try:
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Замены {day.capitalize()}"

        now = datetime.now()
        date_day_str = f"{now:%d.%m.%Y} {day}"

        ws["A1"] = "Замены в расписании"
        ws["A1"].font = Font(bold=True, size=13)
        ws.merge_cells("A1:G1")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 20

        ws["B2"] = date_day_str
        ws["B2"].font = Font(bold=True)
        ws.merge_cells("B2:G2")
        ws["B2"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 18

        ws.row_dimensions[3].height = 6

        ws["A4"] = f"Замены на {now:%d.%m.%Y} ({day.capitalize()})"
        ws["A4"].font = Font(italic=True)
        ws.merge_cells("A4:G4")
        ws["A4"].alignment = Alignment(horizontal="left")
        ws.row_dimensions[4].height = 16

        header_fill   = PatternFill("solid", fgColor="D9E1F2")
        header_font   = Font(bold=True)
        header_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_side     = Side(style="thin")
        border        = Border(left=thin_side, right=thin_side,
                               top=thin_side,  bottom=thin_side)

        headers = [
            ("A5", "№"),
            ("B5", "Урок\n(исходный)"),
            ("C5", "Класс и предмет"),
            ("D5", ""),
            ("E5", "Урок\n(заменяемый)"),
            ("F5", "Замещающий учитель"),
            ("G5", "Примечание"),
        ]
        for cell_addr, title in headers:
            c = ws[cell_addr]
            c.value     = title
            c.font      = header_font
            c.fill      = header_fill
            c.alignment = header_align
            c.border    = border
        ws.row_dimensions[5].height = 30

        col_widths = {"A": 5, "B": 10, "C": 28, "D": 6, "E": 10, "F": 24, "G": 30}
        for col, width in col_widths.items():
            ws.column_dimensions[col].width = width

        data_align  = Alignment(vertical="center", wrap_text=True)
        num_align   = Alignment(horizontal="center", vertical="center")
        odd_fill    = PatternFill("solid", fgColor="EEF2FA")
        even_fill   = PatternFill("solid", fgColor="FFFFFF")

        if temp_rows:
            rows_data = temp_rows
        else:
            rows_data = []
            for cn in sorted(changes.keys()):
                if not re.match(r"^\d{1,2}[А-Я]$", cn):
                    continue
                for ln in sorted(changes[cn].keys(),
                                 key=lambda x: int(x.split("-")[0]) if x.replace("-","").isdigit() else 9999):
                    ct = changes[cn][ln]
                    rt = ""
                    note = ""
                    m = re.search(r"заменён на .+ \((.+?)\)", ct)
                    if m:
                        rt = m.group(1)
                    elif "отменён" in ct:
                        m2 = re.search(r"отменён(?: \((.+?)\))?", ct)
                        note = f"отменён{f', {m2.group(1)}' if m2 and m2.group(1) else ''}"
                    rows_data.append({
                        "ln_orig":    ln,
                        "class_subj": f"{cn} {ct.split(' заменён')[0].split(' отменён')[0]}",
                        "ln_repl":    ln,
                        "replacement": rt,
                        "note":       note,
                    })

        for idx, row in enumerate(rows_data, 1):
            ri   = idx + 5
            fill = odd_fill if idx % 2 == 1 else even_fill

            cells_values = [
                ("A", idx),
                ("B", row.get("ln_orig", "")),
                ("C", row.get("class_subj", "")),
                ("D", ""),
                ("E", row.get("ln_repl", "")),
                ("F", row.get("replacement", "")),
                ("G", row.get("note", "")),
            ]
            for col_letter, val in cells_values:
                c = ws.cell(row=ri, column=ord(col_letter) - ord("A") + 1)
                c.value     = val
                c.fill      = fill
                c.border    = border
                c.alignment = num_align if col_letter in ("A", "B", "E") else data_align
            ws.row_dimensions[ri].height = 18

        wb.save(output_path)
        return True
    except Exception as e:
        await log_error(f"Ошибка создания XLSX: {e}")
        return False


# ─── Общая логика обработки сообщений (VK + MAX) ─────────────────────────────

async def handle_incoming(platform: str, uid: int, text: str, payload: dict, responder: Responder):
    global total_requests, menu_requests, class_requests, teacher_requests

    tlow = (text or "").strip().lower()

    # ── Фиксированные команды (работают в любом состоянии диалога) ───────────
    if tlow in ("начало", "старт", "start", "/start", "помощь", "help"):
        await responder.answer(START_TEXT)
        return

    if tlow in ("учителя", "/teachers"):
        if not teacher_days:
            await responder.answer("Список учителей пуст.")
        else:
            lines = [f"{t}: {', '.join(sorted(d.capitalize() for d in ds))}"
                     for t, ds in sorted(teacher_days.items())]
            await responder.answer("Список учителей:\n" + "\n".join(lines))
        return

    if tlow in ("debug_cache", "/debug_cache"):
        if not changes_cache:
            await responder.answer("Кэш изменений пуст.")
            return
        lines = []
        for day, entry in changes_cache.items():
            date_str = entry.get("date")
            lines.append(f"{day.capitalize()}{f' ({date_str})' if date_str else ''}:")
            for entity, lessons in entry["changes"].items():
                lines.append(f"  {entity}:")
                for ln, val in sorted(lessons.items(), key=_sort_key):
                    lines.append(f"    Урок {ln}: {val}")
        await responder.answer("\n".join(lines))
        return

    # ── Обработка payload (нажатия кнопок) ────────────────────────────────────
    if payload:
        cmd = payload.get("cmd", "")

        if cmd == "changes_day":
            day = payload["day"]
            set_user_state(platform, uid, STATE_CHOOSE_TEACHER,
                           {"change_day": day, "temp_changes": {}, "corrections": []})
            await responder.answer(
                f"Начато редактирование замен на {day.capitalize()}.\n"
                "Введите фамилию педагога или «конец» для завершения."
            )
            return

        if cmd == "menu_schedule":
            resp = "Расписание трапез\n\n"
            for corpus, meals in MEAL_SCHEDULE.items():
                resp += f"{corpus.capitalize()}:\n"
                for m in meals:
                    resp += f"  {m['трапеза'].capitalize()}  {m['время']}  {m['классы']}\n"
                resp += "\n"
            await responder.answer(resp)
            return

        if cmd == "menu_day":
            menu_requests += 1
            day = payload["day"]
            await responder.answer(f"Выберите классы для меню на {day.capitalize()}:",
                                   keyboard=nkb_menu_classes(day))
            return

        if cmd == "menu_classes":
            menu_requests += 1
            day   = payload["day"]
            cl    = payload["cl"]
            cache = menu_1_4_cache if cl == "1-4" else menu_5_11_cache
            vi    = get_current_menu_variant() - 1
            resp  = f"Меню на {day.capitalize()} ({cl} классы):\n\n"
            if day in cache and len(cache[day]) > vi:
                for mtype, dishes in cache[day][vi].items():
                    resp += f"{mtype.capitalize()}:\n"
                    for num, dish in sorted(dishes):
                        resp += f"  {num}. {dish}\n"
                    resp += "\n"
            else:
                resp += "Меню отсутствует.\n"
            await responder.answer(resp)
            return

        if cmd == "lesson":
            st = get_user_state(platform, uid)
            if st["state"] != STATE_CHOOSE_LESSON:
                await responder.answer("Нет активного сеанса. Введите «замены».")
                return
            data = st["data"]
            ln   = str(payload["num"])
            day  = data["change_day"]
            surn = data["current_teacher"]
            entry = next((e for e in teacher_schedule_list
                          if e["surname"].lower() == surn.lower()
                          and e["day"].lower() == day
                          and e["lesson_num"] == ln), None)
            if not entry:
                await responder.answer("Урок не найден.")
                return
            cn, subj = entry["info"].split(" - ", 1)
            busy  = {e["surname"] for e in teacher_schedule_list
                     if e["day"].lower() == day and e["lesson_num"] == ln}
            avail = sorted(t for t in teacher_days if t not in busy)
            data.update(current_lesson=ln, current_class=cn,
                        current_subject=subj, avail_teachers=avail)
            set_user_state(platform, uid, STATE_CHOOSE_REPLACEMENT, data)
            total = len(avail)
            pages = (total + TEACHERS_PER_PAGE - 1) // TEACHERS_PER_PAGE if total else 1
            resp  = (f"Замена для урока {ln} ({cn} – {subj}) на {day.capitalize()}.\n"
                     f"Выберите замену (стр. 1/{pages}):")
            await responder.answer(resp, keyboard=nkb_teachers_page(avail, 0))
            return

        if cmd == "lessons_done":
            st = get_user_state(platform, uid)
            if st["state"] != STATE_CHOOSE_LESSON:
                return
            set_user_state(platform, uid, STATE_CHOOSE_TEACHER, st["data"])
            await responder.answer("Введите следующую фамилию педагога или «конец».")
            return

        if cmd == "repl_page":
            st = get_user_state(platform, uid)
            if st["state"] != STATE_CHOOSE_REPLACEMENT:
                return
            data  = st["data"]
            page  = int(payload.get("page", 0))
            avail = data.get("avail_teachers", [])
            total = len(avail)
            pages = (total + TEACHERS_PER_PAGE - 1) // TEACHERS_PER_PAGE if total else 1
            ln    = data["current_lesson"]
            cn    = data["current_class"]
            subj  = data["current_subject"]
            day   = data["change_day"]
            resp  = (f"Замена для урока {ln} ({cn} – {subj}) на {day.capitalize()}.\n"
                     f"Выберите замену (стр. {page + 1}/{pages}):")
            await responder.answer(resp, keyboard=nkb_teachers_page(avail, page))
            return

        if cmd == "cancel_lesson":
            st = get_user_state(platform, uid)
            if st["state"] != STATE_CHOOSE_REPLACEMENT:
                return
            data = st["data"]
            data["replacement_teacher"] = None
            set_user_state(platform, uid, STATE_ADD_COMMENT, data)
            await responder.answer("Введите комментарий к отмене (или «-» для пропуска):")
            return

        if cmd == "replace":
            st = get_user_state(platform, uid)
            if st["state"] != STATE_CHOOSE_REPLACEMENT:
                return
            data = st["data"]
            rt   = payload.get("teacher", "")
            day  = data["change_day"]
            ln   = data["current_lesson"]
            data["replacement_teacher"] = rt
            set_user_state(platform, uid, STATE_CONFIRM_REPLACEMENT, data)
            lessons = [e for e in teacher_schedule_list
                       if e["surname"].lower() == rt.lower() and e["day"].lower() == day]
            resp = f"Расписание {rt} на {day.capitalize()}:\n"
            if lessons:
                for e in sorted(lessons, key=lambda x: int(x["lesson_num"])):
                    resp += f"  {e['lesson_num']}. {e['info']}\n"
            else:
                resp += "  Уроков нет.\n"
            resp += f"\nНазначить {rt} на урок {ln}?"
            await responder.answer(resp, keyboard=nkb_confirm())
            return

        if cmd == "confirm_yes":
            st = get_user_state(platform, uid)
            if st["state"] != STATE_CONFIRM_REPLACEMENT:
                return
            set_user_state(platform, uid, STATE_ADD_COMMENT, st["data"])
            await responder.answer("Введите комментарий (или «-» для пропуска):")
            return

        if cmd == "confirm_no":
            st = get_user_state(platform, uid)
            if st["state"] != STATE_CONFIRM_REPLACEMENT:
                return
            data  = st["data"]
            avail = data.get("avail_teachers", [])
            ln    = data["current_lesson"]
            cn    = data["current_class"]
            subj  = data["current_subject"]
            day   = data["change_day"]
            total = len(avail)
            pages = (total + TEACHERS_PER_PAGE - 1) // TEACHERS_PER_PAGE if total else 1
            set_user_state(platform, uid, STATE_CHOOSE_REPLACEMENT, data)
            resp  = (f"Замена для урока {ln} ({cn} – {subj}) на {day.capitalize()}.\n"
                     f"Выберите замену (стр. 1/{pages}):")
            await responder.answer(resp, keyboard=nkb_teachers_page(avail, 0))
            return

        if cmd == "confirm_final":
            st = get_user_state(platform, uid)
            if st["state"] != STATE_CONFIRM_CHANGES:
                return
            data = st["data"]
            day  = data.get("change_day")
            temp = data.get("temp_changes", {})
            if day in temp:
                path = os.path.join(CHANGES_DIR,
                                    f"Замены_{day}_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
                temp_rows = data.get("temp_rows", [])
                ok = await generate_xlsx_changes(day, temp[day], path, temp_rows)
                if ok:
                    await responder.answer(f"Замены на {day.capitalize()} сохранены ✅")
                    if not await responder.send_document(path, f"Замены_{day}.xlsx"):
                        await responder.answer("Не удалось отправить файл.")
                    if os.path.exists(path):
                        os.remove(path)
                else:
                    await responder.answer("Ошибка создания файла.")
            else:
                await responder.answer("Замены не были добавлены.")
            clear_user_state(platform, uid)
            return

        if cmd == "cancel_final":
            clear_user_state(platform, uid)
            await responder.answer("Редактирование замен отменено.")
            return

    # ── Текстовые ответы в диалоге ────────────────────────────────────────────

    st = get_user_state(platform, uid)

    if st["state"] == STATE_CHOOSE_TEACHER:
        data = st["data"]
        if tlow == "конец":
            day  = data.get("change_day", "")
            temp = data.get("temp_changes", {})
            resp = f"Замены на {day.capitalize()}:\n\n"
            if day in temp:
                for ent, lessons in temp[day].items():
                    resp += f"{ent}:\n"
                    for ln, v in sorted(lessons.items(), key=_sort_key):
                        resp += f"  Урок {ln}: {v}\n"
            else:
                resp += "Замены не добавлены.\n"
            set_user_state(platform, uid, STATE_CONFIRM_CHANGES, data)
            await responder.answer(resp + "\nПодтвердить сохранение?",
                                   keyboard=nkb_final_confirm())
            return

        matches = get_close_matches(tlow, [t.lower() for t in teacher_days], n=1, cutoff=0.6)
        if not matches:
            await responder.answer(f"Учитель «{text}» не найден. Введите другую фамилию или «конец».")
            return
        surn = next(t for t in teacher_days if t.lower() == matches[0])
        if tlow != surn.lower():
            data.setdefault("corrections", []).append(f"{text.capitalize()} → {surn}")
        data["current_teacher"] = surn
        set_user_state(platform, uid, STATE_CHOOSE_LESSON, data)
        day     = data["change_day"]
        lessons = [e for e in teacher_schedule_list
                   if e["surname"].lower() == surn.lower() and e["day"].lower() == day]
        nums    = sorted({e["lesson_num"] for e in lessons}, key=int)
        resp    = f"Расписание {surn} на {day.capitalize()}:\n"
        if not lessons:
            resp += "Уроков нет.\n"
            await responder.answer(resp + "Введите другую фамилию или «конец».")
            set_user_state(platform, uid, STATE_CHOOSE_TEACHER, data)
            return
        for e in sorted(lessons, key=lambda x: int(x["lesson_num"])):
            resp += f"  {e['lesson_num']}. {e['info']}\n"
        await responder.answer(resp + "\nВыберите урок:", keyboard=nkb_lessons(nums))
        return

    if st["state"] in (STATE_CHOOSE_REPLACEMENT, STATE_CONFIRM_REPLACEMENT):
        await responder.answer("Используйте кнопки для выбора.")
        return

    if st["state"] == STATE_ADD_COMMENT:
        data    = st["data"]
        comment = "" if tlow == "-" else text
        day     = data["change_day"]
        ln      = data["current_lesson"]
        cn      = data["current_class"]
        subj    = data["current_subject"]
        surn    = data["current_teacher"]
        rt      = data.get("replacement_teacher")
        data["temp_changes"].setdefault(day, {})
        data.setdefault("temp_rows", [])
        if rt:
            rs = comment if comment else subj
            ct = f"{subj} заменён на {rs} ({rt})"
            if comment and comment != subj:
                ct += f" ({comment})"
            data["temp_changes"][day].setdefault(cn, {})[ln] = ct
            data["temp_changes"][day].setdefault(rt, {})[ln] = f"{cn} - {ct}"
            data["temp_rows"].append({
                "ln_orig":   ln,
                "class_subj": f"{cn} {subj}",
                "ln_repl":   ln,
                "replacement": rt,
                "note":      comment,
            })
        else:
            ct = f"{subj} отменён" + (f" ({comment})" if comment else "")
            data["temp_changes"][day].setdefault(cn, {})[ln] = ct
            note_text = f"отменён{f', {comment}' if comment else ''}"
            data["temp_rows"].append({
                "ln_orig":   ln,
                "class_subj": f"{cn} {subj}",
                "ln_repl":   ln,
                "replacement": "",
                "note":      note_text,
            })
        set_user_state(platform, uid, STATE_CHOOSE_LESSON, data)
        lessons = [e for e in teacher_schedule_list
                   if e["surname"].lower() == surn.lower() and e["day"].lower() == day]
        nums = sorted({e["lesson_num"] for e in lessons}, key=int)
        resp = f"Добавлено ✅\nРасписание {surn} на {day.capitalize()}:\n"
        for e in sorted(lessons, key=lambda x: int(x["lesson_num"])):
            resp += f"  {e['lesson_num']}. {e['info']}\n"
        await responder.answer(resp + "\nВыберите урок:", keyboard=nkb_lessons(nums))
        return

    if st["state"] == STATE_CONFIRM_CHANGES:
        await responder.answer("Используйте кнопки для подтверждения.")
        return

    # ── Обычные запросы ───────────────────────────────────────────────────────

    if tlow == "меню":
        menu_requests += 1
        await responder.answer("Выберите день или расписание трапез:", keyboard=nkb_menu_main())
        return

    if tlow == "замены":
        total_requests += 1
        await responder.answer("Выберите день для ввода замен:", keyboard=nkb_days_changes())
        return

    words = tlow.split()
    if not words:
        await responder.answer("Введите запрос: «9А понедельник» или «Иванов вторник».")
        return

    day = None
    for w in words:
        if w in DAY_MAPPING:
            day = DAY_MAPPING[w]; break
        if w in DAYS:
            day = w; break
        cm = get_close_matches(w, DAYS, n=1, cutoff=0.6)
        if cm:
            day = cm[0]; break

    if not day:
        await log_error(f"День не распознан: {text}")
        await responder.answer("День недели не распознан. Пример: «понедельник», «вт.»")
        return

    entity = " ".join(w for w in words if w not in DAY_MAPPING and w not in DAYS).strip()
    if not entity:
        await responder.answer("Укажите класс (например «9А») или фамилию учителя.")
        return

    total_requests += 1
    ecl = "".join(entity.split()).upper()

    if re.match(r"^\d{1,2}[А-Я]$", ecl):
        class_requests[ecl] = class_requests.get(ecl, 0) + 1
        resp = f"Расписание {ecl} на {day.capitalize()}:\n\n"
        if ecl in schedule_by_class and day in schedule_by_class[ecl]:
            for ln, ls in sorted(schedule_by_class[ecl][day].items(), key=lambda x: int(x[0])):
                resp += f"{ln}. {', '.join(ls)}\n"
        else:
            resp += "Уроки не найдены.\n"
        if ecl in consultations_by_class and day in consultations_by_class[ecl]:
            resp += "\nКонсультации:\n"
            for ln, ls in sorted(consultations_by_class[ecl][day].items(), key=lambda x: int(x[0])):
                resp += f"{ln}. {', '.join(ls)}\n"
        resp += await get_changes_for_class(ecl, day)
        await responder.answer(resp)
        return

    matches = get_close_matches(entity, [t.lower() for t in teacher_days], n=1, cutoff=0.6)
    if matches:
        surn = next(t for t in teacher_days if t.lower() == matches[0])
        if entity != surn.lower():
            await responder.answer(f"Исправлено: {entity.capitalize()} → {surn}")
        teacher_requests[surn] = teacher_requests.get(surn, 0) + 1
        resp = f"Расписание {surn} на {day.capitalize()}:\n\n"
        tl = [e for e in teacher_schedule_list
              if e["surname"].lower() == surn.lower() and e["day"].lower() == day]
        if tl:
            for e in sorted(tl, key=lambda x: int(x["lesson_num"])):
                resp += f"{e['lesson_num']}. {e['info']}\n"
        else:
            resp += "Уроки не найдены.\n"
        if surn in consultations_by_teacher and day in consultations_by_teacher[surn]:
            resp += "\nКонсультации:\n"
            for ln, ls in sorted(consultations_by_teacher[surn][day].items(), key=lambda x: int(x[0])):
                resp += f"{ln}. {', '.join(ls)}\n"
        resp += await get_changes_for_teacher(surn, day)
        await responder.answer(resp)
        return

    if "меню" in tlow:
        menu_requests += 1
        await responder.answer(f"Выберите классы для меню на {day.capitalize()}:",
                               keyboard=nkb_menu_classes(day))
        return

    if "изменения" in tlow:
        if day in changes_cache:
            entry = changes_cache[day]
            date_str = entry.get("date")
            resp = f"Изменения{f' на {date_str}' if date_str else ''} ({day.capitalize()}):\n"
            for ent, lessons in entry["changes"].items():
                resp += f"\n{ent}:\n"
                for ln, v in sorted(lessons.items(), key=_sort_key):
                    resp += f"  Урок {ln}: {v}\n"
        else:
            resp = f"Изменений на {day.capitalize()} нет."
        await responder.answer(resp)
        return

    await log_error(f"Не найдено: '{entity}' | запрос: {text}")
    await responder.answer(f"Класс или учитель «{entity}» не найден. Попробуйте ещё раз.")


# ─── VK: обработчики ──────────────────────────────────────────────────────────

@bot.on.message(attachment=["doc"])
async def vk_handle_doc_upload(message: Message):
    responder = VKResponder(message)
    for att in (message.attachments or []):
        doc = att.doc
        if not doc:
            continue
        fname = doc.title or ""

        async def _get_bytes(url=doc.url):
            async with aiohttp.ClientSession() as session:
                async with session.get(url) as resp:
                    return await resp.read()

        await handle_incoming_file(_get_bytes(), fname, responder)


@bot.on.message()
async def vk_handle_message(message: Message):
    responder = VKResponder(message)
    payload = parse_payload(message)
    text = message.text or ""
    await handle_incoming("vk", message.from_id, text, payload, responder)


# ─── MAX: long-polling клиент ─────────────────────────────────────────────────

async def max_handle_update(upd: dict):
    try:
        utype = upd.get("update_type")

        if utype == "message_created":
            message = upd.get("message") or {}
            sender = message.get("sender") or {}
            uid = sender.get("user_id")
            if uid is None:
                return
            recipient = message.get("recipient") or {}
            chat_id = recipient.get("chat_id")
            is_chat = recipient.get("chat_type") in ("chat", "channel") and chat_id is not None

            body = message.get("body") or {}
            attachments = body.get("attachments") or []
            file_att = next((a for a in attachments if a.get("type") == "file"), None)

            responder = MAXResponder(uid, chat_id=chat_id if is_chat else None)

            if file_att:
                fpayload = file_att.get("payload") or {}
                url = fpayload.get("url")
                fname = fpayload.get("filename") or fpayload.get("name") or "file"
                if url:
                    async def _get_bytes(u=url):
                        async with max_client_session() as session:
                            async with session.get(u, headers={"Authorization": MAX_TOKEN}) as resp:
                                return await resp.read()
                    await handle_incoming_file(_get_bytes(), fname, responder)
                    return

            text = body.get("text") or ""
            await handle_incoming("max", uid, text, {}, responder)

        elif utype == "message_callback":
            callback = upd.get("callback") or {}
            user = callback.get("user") or {}
            uid = user.get("user_id")
            if uid is None:
                return
            payload_raw = callback.get("payload") or ""
            try:
                payload = json.loads(payload_raw) if payload_raw else {}
            except Exception:
                payload = {}

            message = upd.get("message") or {}
            recipient = message.get("recipient") or {}
            chat_id = recipient.get("chat_id")
            is_chat = recipient.get("chat_type") in ("chat", "channel") and chat_id is not None

            responder = MAXResponder(uid, chat_id=chat_id if is_chat else None,
                                     callback_id=callback.get("callback_id"))
            await handle_incoming("max", uid, "", payload, responder)

        elif utype == "bot_started":
            user = upd.get("user") or {}
            uid = user.get("user_id")
            if uid is not None:
                await max_send_message(START_TEXT, user_id=uid)

        # message_edited / message_removed / прочие события — не обрабатываем

    except Exception as e:
        await log_error(f"MAX: ошибка обработки события: {e}")


async def max_polling_loop():
    await ensure_ru_trusted_ca()
    marker = None
    async with max_client_session() as session:
        while True:
            try:
                params = {"timeout": 30, "limit": 100}
                if marker is not None:
                    params["marker"] = marker
                status, data = await max_api_request(session, "GET", "/updates", params=params)
                if status != 200:
                    if status == 401:
                        await log_error("MAX: неверный токен доступа (401), опрос остановлен на 60с")
                        await asyncio.sleep(60)
                    else:
                        await log_error(f"MAX getUpdates ошибка {status}: {data}")
                        await asyncio.sleep(5)
                    continue
                for upd in data.get("updates", []):
                    asyncio.ensure_future(max_handle_update(upd))
                if data.get("marker") is not None:
                    marker = data["marker"]
            except Exception as e:
                await log_error(f"MAX: исключение в цикле опроса: {e}")
                await asyncio.sleep(5)
    print(f"{get_timestamp()} - MAX бот запущен ✅ (long polling)")


# ─── Запуск ───────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    print(f"{get_timestamp()} - Загрузка расписания...")
    load_or_parse_schedule()
    print(f"{get_timestamp()} - Загрузка меню...")
    load_or_parse_menu()

    if os.path.exists(CHANGES_CACHE_FILE):
        with open(CHANGES_CACHE_FILE, "rb") as f:
            changes_cache = pickle.load(f)
        needs_rebuild = False
        for day_key, val in changes_cache.items():
            if not isinstance(val, dict) or "changes" not in val:
                needs_rebuild = True
                break
        if needs_rebuild:
            print(f"{get_timestamp()} - Старый формат кэша изменений, пересборка...")
            changes_cache = {}
            os.remove(CHANGES_CACHE_FILE)
        else:
            print(f"{get_timestamp()} - Кэш изменений загружен")

    bot.startup_tasks.append(switch_menu_variant())
    bot.startup_tasks.append(check_midnight())
    # MAX работает в том же event loop, что и VK — общие кэши и состояния
    # доступны обоим ботам без какой-либо синхронизации между процессами.
    bot.startup_tasks.append(max_polling_loop())

    print(f"{get_timestamp()} - VK бот запущен ✅")
    print(f"{get_timestamp()} - Запуск MAX бота (long polling)...")
    bot.run()
